from __future__ import annotations

import base64
import html
import io
import json
import re
import zipfile
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .auth import AuthenticatedUser
from .research_versions import ResearchVersionStore, SECTION_LABELS


def _text(section: Any) -> str:
    if isinstance(section, dict):
        return str(section.get("text") or "").strip()
    return str(section or "").strip()


def _sources(document: dict[str, Any]) -> list[dict[str, Any]]:
    rows = ((document.get("evidence_manifest") or {}).get("sources") or [])
    return sorted(rows, key=lambda row: int(row.get("citation_number") or 999999))


def _authors(source: dict[str, Any]) -> list[str]:
    value = source.get("authors") or []
    if isinstance(value, str):
        return [value]
    out: list[str] = []
    for item in value:
        if isinstance(item, dict):
            name = item.get("name") or item.get("display_name")
        else:
            name = item
        if name:
            out.append(str(name))
    return out


def _reference(source: dict[str, Any]) -> str:
    value = str(source.get("reference") or source.get("title") or "Fuente").strip()
    return value


def _ris(document: dict[str, Any]) -> str:
    entries: list[str] = []
    for source in _sources(document):
        lines = ["TY  - JOUR"]
        for author in _authors(source):
            lines.append(f"AU  - {author}")
        if source.get("title"):
            lines.append(f"TI  - {source['title']}")
        if source.get("journal"):
            lines.append(f"JO  - {source['journal']}")
        year = str(source.get("publication_date") or "")[:4]
        if year:
            lines.append(f"PY  - {year}")
        if source.get("doi"):
            lines.append(f"DO  - {source['doi']}")
        url = source.get("source_url") or source.get("access_url")
        if url:
            lines.append(f"UR  - {url}")
        lines.append("ER  - ")
        entries.append("\n".join(lines))
    return "\n\n".join(entries) + ("\n" if entries else "")


def _docx(document: dict[str, Any]) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    title = str(document.get("topic") or "ScientificBrain Research")
    p = doc.add_paragraph()
    p.alignment = 1
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    brief = document.get("research_brief") or {}
    problem = str(brief.get("problem_statement") or "").strip()
    if problem:
        doc.add_heading("Problema científico", level=2)
        doc.add_paragraph(problem)
    for key, label in SECTION_LABELS.items():
        value = _text((document.get("sections") or {}).get(key))
        if not value:
            continue
        doc.add_heading(label, level=1)
        for block in re.split(r"\n\s*\n", value):
            if block.strip():
                doc.add_paragraph(block.strip())
    sources = _sources(document)
    if sources:
        doc.add_heading("Referencias", level=1)
        for source in sources:
            doc.add_paragraph(_reference(source))
    props = doc.core_properties
    props.title = title
    props.subject = "ScientificBrain evidence-grounded research export"
    props.comments = "Generated from a versioned ScientificBrain research document."
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf(document: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("SBTitle", parent=styles["Title"], alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=16, leading=19, spaceAfter=12)
    head_style = ParagraphStyle("SBHead", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=14, spaceBefore=10, spaceAfter=5)
    body_style = ParagraphStyle("SBBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=7)
    small_style = ParagraphStyle("SBSmall", parent=body_style, fontSize=8.5, leading=11)
    doc = SimpleDocTemplate(buf, pagesize=LETTER, rightMargin=0.65 * inch, leftMargin=0.65 * inch, topMargin=0.65 * inch, bottomMargin=0.65 * inch, title=str(document.get("topic") or "ScientificBrain Research"), author="ScientificBrain")
    story: list[Any] = [Paragraph(html.escape(str(document.get("topic") or "ScientificBrain Research")), title_style)]
    brief = document.get("research_brief") or {}
    problem = str(brief.get("problem_statement") or "").strip()
    if problem:
        story += [Paragraph("Problema científico", head_style), Paragraph(html.escape(problem).replace("\n", "<br/>"), body_style)]
    for key, label in SECTION_LABELS.items():
        value = _text((document.get("sections") or {}).get(key))
        if not value:
            continue
        story.append(Paragraph(html.escape(label), head_style))
        for block in re.split(r"\n\s*\n", value):
            if block.strip():
                story.append(Paragraph(html.escape(block.strip()).replace("\n", "<br/>"), body_style))
    sources = _sources(document)
    if sources:
        story += [Spacer(1, 6), Paragraph("Referencias", head_style)]
        for source in sources:
            story.append(Paragraph(html.escape(_reference(source)), small_style))
    doc.build(story)
    return buf.getvalue()


def _binary_response(filename: str, mime: str, data: bytes, fmt: str) -> dict[str, Any]:
    return {
        "format": fmt,
        "mime": mime,
        "filename": filename,
        "encoding": "base64",
        "content_base64": base64.b64encode(data).decode("ascii"),
        "byte_length": len(data),
    }


@dataclass
class ProfessionalExportService:
    user: AuthenticatedUser
    folder_id: str

    def __post_init__(self) -> None:
        self.store = ResearchVersionStore(self.user, self.folder_id)

    def export(self, fmt: str) -> dict[str, Any]:
        fmt = fmt.strip().lower()
        document = self.store.document()
        if not document:
            raise KeyError("research_document_not_found")
        title = re.sub(r"[^A-Za-z0-9._-]+", "_", str(document.get("topic") or "ScientificBrain_Research")).strip("._")[:90] or "ScientificBrain_Research"
        if fmt in {"markdown", "md", "latex", "tex", "bibtex", "bib", "json", "scibrain"}:
            return self.store.export(fmt)
        if fmt == "ris":
            return {"format": "ris", "mime": "application/x-research-info-systems", "filename": f"{title}.ris", "content": _ris(document), "encoding": "utf-8"}
        if fmt in {"docx", "word"}:
            return _binary_response(f"{title}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _docx(document), "docx")
        if fmt == "pdf":
            return _binary_response(f"{title}.pdf", "application/pdf", _pdf(document), "pdf")
        if fmt in {"package", "zip"}:
            memory = io.BytesIO()
            with zipfile.ZipFile(memory, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                md = self.store.export("markdown")
                bib = self.store.export("bibtex")
                tex = self.store.export("latex")
                zf.writestr("research.md", md.get("content") or "")
                zf.writestr("references.bib", bib.get("content") or "")
                zf.writestr("research.tex", tex.get("content") or "")
                zf.writestr("references.ris", _ris(document))
                zf.writestr("scientificbrain.json", json.dumps(document, ensure_ascii=False, indent=2, default=str))
            return _binary_response(f"{title}_reproducible.zip", "application/zip", memory.getvalue(), "package")
        raise ValueError(f"Unsupported export format: {fmt}")
